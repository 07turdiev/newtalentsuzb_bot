"""Arizalarni Excel jadval va ZIP arxiv ko'rinishida yuklab olish."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .db import Application, Contest
from .forms import ALL_FIELDS, display
from .locales import t

COLUMNS = (
    ("№", 7),
    ("Sana", 17),
    ("Turi", 14),
    ("F.I.Sh. / Jamoa nomi", 32),
    ("Ishtirokchilar", 40),
    ("Soni", 7),
    ("Tug'ilgan sana", 13),
    ("Ta'lim muassasasi", 32),
    ("Yo'nalish", 20),
    ("Yosh toifasi", 14),
    ("Asar nomi va muallifi", 34),
    ("O'qituvchi", 26),
    ("Konsertmeyster", 26),
    ("Ota-onasi", 30),
    ("Muassasa rahbari", 26),
    ("Telefon", 18),
    ("Ijro videosi", 30),
    ("To'lov cheki", 17),
    ("Holat", 24),
    ("Telegram ID", 14),
)
KIND_SHORT = {"solo": "Yakka", "team": "Jamoaviy"}


def _row(app: Application) -> list:
    d = app.data

    def show(key: str) -> str:
        return display(ALL_FIELDS[key], d.get(key))

    participants = d.get("participants") or []
    return [
        app.id,
        app.created_at[:16],
        KIND_SHORT.get(app.kind, app.kind),
        app.display_name,
        show("participants") if participants else "",
        len(participants) if participants else 1,
        d.get("birth_date") or "",
        show("institution"),
        show("direction"),
        show("age_category"),
        show("piece"),
        show("teacher"),
        show("accompanist"),
        show("parents"),
        show("head"),
        show("phone"),
        show("video") if d.get("video") else "",
        app.receipt["at"][:16] if app.receipt else "",
        t("uz", f"status.{app.status}").split(" ", 1)[-1],
        app.user_id,
    ]


def build_excel(contest: Contest, apps: list[Application]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Arizalar"

    ws.append([name for name, _ in COLUMNS])
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell, (_, width) in zip(ws[1], COLUMNS):
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for app in apps:
        ws.append(_row(app))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_zip(apps: list[Application], data_dir: Path) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for app in apps:
            if not app.doc_path:
                continue
            path = data_dir / app.doc_path
            if path.exists():
                archive.write(path, arcname=path.name)
    return buffer.getvalue()
