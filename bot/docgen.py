"""templates/ dagi Word ariza shakllarini to'ldirish.

Shablonlardagi har bir qator yorlig'i (masalan «Ўқитувчиси Ф.И.О») bo'yicha topiladi va
undagi «_____» chizig'i o'rniga javob yoziladi. Shablon matnini o'zgartirsangiz, quyidagi
LABELS ro'yxatidagi yorliqlar bilan mos kelishini tekshiring.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .config import TEMPLATES_DIR
from .forms import format_phone

TEMPLATES = {
    "solo": TEMPLATES_DIR / "solo.docx",
    "team": TEMPLATES_DIR / "team.docx",
}

# (qator boshidagi yorliq, maydon kaliti)
LABELS = (
    ("Жамоа ва ансамбль номи", "team_name"),
    ("Ҳар бир иштирокчи", "participants"),
    ("Жамоа ва ансамблдаги иштирокчилар жами сони", "participants_count"),
    ("Иштирокчининг Ф.И.О", "full_name"),
    ("Туғилган", "birth_date"),
    ("Таҳсил оладиган", "institution"),
    ("Иштирок этадиган йўналиши", "direction"),
    ("Иштирок этадиган ёш", "age_category"),
    ("Ижро этадиган асар", "piece"),
    ("Ўқитувчиси", "teacher"),
    ("Концертмейстери", "accompanist"),
    ("Ота ва онаси", "parents"),
    ("Таълим муассаса раҳбари", "head"),
    ("Иштирокчининг телефон", "phone"),
    ("Иштирокчининг розилик", "consent"),
)
TITLE_PLACEHOLDER = "Танлов номи"
EMPTY = "—"
LINE = re.compile(r"_+")


def _match_label(text: str) -> str | None:
    text = text.strip()
    for label, key in LABELS:
        if text.startswith(label):
            return key
    return None


def _is_line_only(text: str) -> bool:
    """Faqat «____» dan iborat davom qatori."""
    text = text.strip()
    return bool(text) and not text.strip("_")


def _remove(paragraph: Paragraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


def _fill(paragraph: Paragraph, value: str) -> None:
    """Qatordagi birinchi «____» o'rniga qiymatni yozadi, qolgan chiziqlarni o'chiradi."""
    line_runs = [r for r in paragraph.runs if "_" in r.text]
    if not line_runs:
        paragraph.add_run(" " + value).bold = True
        return

    first, *rest = line_runs
    for run in rest:
        run.text = LINE.sub("", run.text)

    match = LINE.search(first.text)
    before = first.text[: match.start()]
    if before and not before[-1].isspace():
        before += " "
    elif not before:
        # Chiziq alohida run'da: oldingi run bo'sh joy bilan tugamasa, bo'sh joy qo'shamiz
        runs = paragraph.runs
        idx = next(i for i, r in enumerate(runs) if r._r is first._r)
        if idx > 0 and runs[idx - 1].text and not runs[idx - 1].text[-1].isspace():
            before = " "
    first.text = before

    value_element = copy.deepcopy(first._r)
    first._r.addnext(value_element)
    value_run = Run(value_element, paragraph)
    value_run.text = value
    value_run.bold = True
    value_run.underline = True


def _fill_list(label: Paragraph, continuation: list[Paragraph], items: list[str]) -> None:
    """Ishtirokchilar ro'yxati: birinchisi yorliq qatoriga, qolganlari alohida qatorlarga."""
    items = items or [EMPTY]
    _fill(label, items[0])
    if not continuation:
        for item in items[1:]:
            run = label.add_run()
            run.add_break()
            run.text = item
        return

    # Birinchi davom qatorini namuna sifatida olib, kerakli miqdorda nusxalaymiz
    template = copy.deepcopy(continuation[0]._p)
    for paragraph in continuation:
        _remove(paragraph)
    anchor = label
    for item in items[1:]:
        element = copy.deepcopy(template)
        anchor._p.addnext(element)
        anchor = Paragraph(element, label._parent)
        _fill(anchor, item)


def _values(kind: str, data: dict[str, Any]) -> dict[str, Any]:
    values: dict[str, Any] = {
        key: str(data.get(key) or EMPTY)
        for _, key in LABELS
        if key not in ("participants", "participants_count", "phone", "consent")
    }
    participants = data.get("participants") or []
    values["participants"] = [f"{i}. {name}" for i, name in enumerate(participants, 1)]
    values["participants_count"] = str(len(participants)) if participants else EMPTY

    phone = data.get("phone") or ""
    if phone.startswith("998") and len(phone) == 12:
        values["phone"] = format_phone(phone).removeprefix("+998 ")
    else:
        values["phone"] = format_phone(phone) if phone else EMPTY

    consent_at = data.get("consent_at")
    values["consent"] = f"Розилик берилган ({consent_at}, Telegram бот орқали)" if consent_at else EMPTY
    return values


def render(kind: str, contest_title: str, data: dict[str, Any], out_path: Path) -> Path:
    doc = Document(str(TEMPLATES[kind]))
    values = _values(kind, data)
    paragraphs = list(doc.paragraphs)

    title = contest_title.strip().strip("\"'“”«»„")
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.text == TITLE_PLACEHOLDER:
                run.text = title

    i = 0
    while i < len(paragraphs):
        paragraph = paragraphs[i]
        key = _match_label(paragraph.text)
        if key is None:
            i += 1
            continue

        j = i + 1
        continuation = []
        while j < len(paragraphs) and _is_line_only(paragraphs[j].text):
            continuation.append(paragraphs[j])
            j += 1

        if key == "participants":
            _fill_list(paragraph, continuation, values["participants"])
        else:
            if key == "phone" and not values["phone"][0].isdigit():
                # Xorijiy raqam: shablondagi «+998» ni olib tashlaymiz
                for run in paragraph.runs:
                    run.text = run.text.replace("+998", "")
            _fill(paragraph, values[key])
            for extra in continuation:
                _remove(extra)
        i = j

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def file_name(app_id: int, display_name: str) -> str:
    slug = re.sub(r"[^\w]+", "_", display_name, flags=re.UNICODE).strip("_")[:40]
    return f"Ariza_{app_id}_{slug}.docx" if slug else f"Ariza_{app_id}.docx"
